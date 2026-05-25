"""
API Test Case Generator — Single sheet có cột "Data Test" + MIX 2 style.

Khác với generate_testcase_xlsx.py (UI/general):
- Thêm cột D "Data Test" sau Pre-Condition (cho API testing)
- Hỗ trợ field-by-field validation với sub-section header per field
- Auto-detect style theo section:
  + Section có "field_groups" → render field validation style (field header vàng nhạt)
  + Section có "test_cases" trực tiếp → render plain style
- 1 sheet có thể mix cả 2 style

Cấu trúc cột:
A: ID | B: Test Item | C: Pre-Condition | D: Data Test | E: Step | F: Expected | G: Priority | H-V: Round 1-3

Cách dùng:

    from generate_api_testcase_xlsx import create_mixed_workbook
    
    sections = [
        # Style 1 — field validation (có sub-section per field)
        {
            "name": "I. API X — Validate tung field",
            "field_groups": [
                {
                    "field": "request_id",
                    "required": "required, String(100)",
                    "test_cases": [
                        {
                            "test_item": "KT de trong truong request_id",
                            "pre_condition": "Da co JWT hop le",
                            "data_test": 'request_id = ""',
                            "step": "1. Cac field khac hop le\\n2. request_id rong\\n3. Send",
                            "expected": "Response: code = 422",
                            "priority": "High",
                        },
                    ]
                },
            ]
        },
        # Style 2 — plain (flow test, security test)
        {
            "name": "II. API X — Flow chinh",
            "test_cases": [
                {"test_item": "Verify happy path", "pre_condition": "...",
                 "data_test": "POST /api/x\\nBody: {...}",
                 "step": "1. Send\\n2. Verify", "expected": "code = 200", "priority": "High"},
            ]
        }
    ]
    
    create_mixed_workbook(
        output_path="/mnt/user-data/outputs/API_TestCases.xlsx",
        module_name="API X",
        module_code="MOD_API",
        creator="QA Team",
        sections=sections,
    )

Dung kem voi:
- references/api-field-validation-pattern.md (standard checklist per field)
- references/security-injection-tests.md (SQL injection deep)
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict

COLOR_MODULE_HEADER = "B4C6E7"
COLOR_ROUND1 = "5BC95E"
COLOR_ROUND2 = "F7CAAC"
COLOR_ROUND3 = "D8D8D8"
COLOR_PASS = "C2D69B"
COLOR_FAIL = "FBD4B4"
COLOR_IMPACT = "B2A1C7"
COLOR_NOTRUN = "BFBFBF"
COLOR_TOTAL = "B8CCE4"
COLOR_MAIN_HEADER = "FFCC66"
COLOR_SECTION = "CFE2F3"
COLOR_FIELD_HEADER = "FFF2CC"
COLOR_PRIORITY_HIGH = "F4CCCC"
COLOR_PRIORITY_MEDIUM = "FFF2CC"
COLOR_PRIORITY_LOW = "D9EAD3"

THIN_BORDER = Border(
    left=Side(style="thin", color="999999"),
    right=Side(style="thin", color="999999"),
    top=Side(style="thin", color="999999"),
    bottom=Side(style="thin", color="999999"),
)

FONT_DEFAULT = Font(name="Arial", size=10)
FONT_BOLD = Font(name="Arial", size=10, bold=True)
FONT_HEADER = Font(name="Arial", size=11, bold=True)
LAST_COL = 22


def _fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def _priority_fill(p):
    p = (p or "").strip().lower()
    if p in ("high", "p1", "critical", "cao"):
        return _fill(COLOR_PRIORITY_HIGH)
    if p in ("medium", "p2", "trung bình", "tb"):
        return _fill(COLOR_PRIORITY_MEDIUM)
    if p in ("low", "p3", "p4", "thấp", "thap"):
        return _fill(COLOR_PRIORITY_LOW)
    return PatternFill()


def _set_header_block(ws, module_code, creator=""):
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 21
    ws.row_dimensions[4].height = 34
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 22
    ws.row_dimensions[7].height = 22

    ws["A2"] = "Module Code"
    ws["A2"].fill = _fill(COLOR_MODULE_HEADER)
    ws["A2"].font = FONT_BOLD
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].border = THIN_BORDER
    ws["B2"] = module_code
    ws["B2"].font = FONT_BOLD
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B2"].border = THIN_BORDER

    for rng, label, color in [("H2:L2", "Round 1", COLOR_ROUND1),
                              ("M2:Q2", "Round 2", COLOR_ROUND2),
                              ("R2:V2", "Round 3", COLOR_ROUND3)]:
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = label
        c.fill = _fill(color)
        c.font = FONT_BOLD
        c.alignment = Alignment(horizontal="center", vertical="center")

    ws["A3"] = "Total of Test cases Created"
    ws["A3"].fill = _fill(COLOR_MODULE_HEADER)
    ws["A3"].font = FONT_BOLD
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A3"].border = THIN_BORDER

    ws["B3"] = '=COUNTIF(F8:F1865,"<>")'
    ws["B3"].font = FONT_BOLD
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B3"].border = THIN_BORDER

    ws["C3"] = "% Executed round 1"
    ws["C3"].fill = _fill(COLOR_MODULE_HEADER)
    ws["C3"].font = FONT_BOLD
    ws["C3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["C3"].border = THIN_BORDER

    ws["D3"] = "=IFERROR(B3/B4,0)"
    ws["D3"].number_format = "0.00%"
    ws["D3"].font = FONT_BOLD
    ws["D3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["D3"].border = THIN_BORDER

    sub_headers = ["Pass", "Fail", "Impact", "Not Run", "Total"]
    sub_fills = [COLOR_PASS, COLOR_FAIL, COLOR_IMPACT, COLOR_NOTRUN, COLOR_TOTAL]
    for start_col in [8, 13, 18]:
        for i, (label, color) in enumerate(zip(sub_headers, sub_fills)):
            cell = ws.cell(row=3, column=start_col + i, value=label)
            cell.fill = _fill(color)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
            cell.font = FONT_DEFAULT

    ws["A4"] = "Total of Test cases plan to Executed"
    ws["A4"].fill = _fill(COLOR_MODULE_HEADER)
    ws["A4"].font = FONT_BOLD
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A4"].border = THIN_BORDER

    ws["B4"] = "=L4+Q4"
    ws["B4"].font = FONT_BOLD
    ws["B4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B4"].border = THIN_BORDER

    ws["C4"] = "Creator"
    ws["C4"].fill = _fill(COLOR_MODULE_HEADER)
    ws["C4"].font = FONT_BOLD
    ws["C4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["C4"].border = THIN_BORDER

    ws["D4"] = creator
    ws["D4"].font = FONT_BOLD
    ws["D4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["D4"].border = THIN_BORDER

    for start_col in [8, 13, 18]:
        col_pass = get_column_letter(start_col)
        col_fail = get_column_letter(start_col + 1)
        col_impact = get_column_letter(start_col + 2)
        col_notrun = get_column_letter(start_col + 3)
        col_total = get_column_letter(start_col + 4)
        ws[f"{col_pass}4"] = f'=COUNTIFS(${col_pass}$8:${col_pass}$1406,"Pass")'
        ws[f"{col_fail}4"] = f'=COUNTIFS(${col_pass}$8:${col_pass}$1406,"Fail")'
        ws[f"{col_impact}4"] = f'=COUNTIFS(${col_pass}$8:${col_pass}$1406,"Impact")'
        ws[f"{col_notrun}4"] = f'=COUNTIFS(${col_pass}$8:${col_pass}$1406,"Not Run")'
        ws[f"{col_total}4"] = f"=SUM({col_pass}4:{col_notrun}4)"
        for cl in [col_pass, col_fail, col_impact, col_notrun, col_total]:
            cell = ws[f"{cl}4"]
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
            cell.font = FONT_DEFAULT

    main_headers = [("A", "ID"), ("B", "Test Item"), ("C", "Pre - Condition"),
                    ("D", "Data Test"), ("E", "Step"), ("F", "Expected Output"),
                    ("G", "Priority")]
    for col, text in main_headers:
        ws.merge_cells(f"{col}6:{col}7")
        cell = ws[f"{col}6"]
        cell.value = text
        cell.fill = _fill(COLOR_MAIN_HEADER)
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for rng, label, color in [("H6:L6", "Round 1", COLOR_ROUND1),
                              ("M6:Q6", "Round 2", COLOR_ROUND2),
                              ("R6:V6", "Round 3", COLOR_ROUND3)]:
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = label
        c.fill = _fill(color)
        c.font = FONT_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")

    round_sub = ["Result", "Bug ID", "Tester", "Test Date", "Note"]
    for start_col, color in [(8, COLOR_ROUND1), (13, COLOR_ROUND2), (18, COLOR_ROUND3)]:
        for i, label in enumerate(round_sub):
            cell = ws.cell(row=7, column=start_col + i, value=label)
            cell.fill = _fill(color)
            cell.font = FONT_BOLD
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER


def _set_column_widths(ws):
    widths = {
        "A": 14, "B": 38, "C": 25, "D": 38, "E": 50, "F": 38, "G": 10,
        "H": 10, "I": 12, "J": 12, "K": 13, "L": 25,
        "M": 10, "N": 12, "O": 12, "P": 13, "Q": 25,
        "R": 10, "S": 12, "T": 12, "U": 13, "V": 25,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _add_section_header(ws, row, name):
    for col_idx in range(1, LAST_COL + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.fill = _fill(COLOR_SECTION)
        cell.font = FONT_BOLD
        if col_idx == 1:
            cell.value = name
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22


def _add_field_header(ws, row, field, required):
    for col_idx in range(1, LAST_COL + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.fill = _fill(COLOR_FIELD_HEADER)
        cell.border = THIN_BORDER
        cell.font = FONT_BOLD
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row, column=1, value=field)
    ws.cell(row=row, column=2, value=required)
    ws.row_dimensions[row].height = 20


def _add_test_case_row(ws, row, tc):
    a_formula = (
        f'=IF(OR(E{row}<>"",F{row}<>""),'
        f'"["&$B$2&"-"&TEXT(ROW()-7-COUNTBLANK($F$8:F{row}),"##")&"]","")'
    )
    ws.cell(row=row, column=1, value=a_formula)
    ws.cell(row=row, column=2, value=tc.get("test_item", ""))
    ws.cell(row=row, column=3, value=tc.get("pre_condition", ""))
    ws.cell(row=row, column=4, value=tc.get("data_test", ""))
    ws.cell(row=row, column=5, value=tc.get("step", ""))
    ws.cell(row=row, column=6, value=tc.get("expected", ""))
    ws.cell(row=row, column=7, value=tc.get("priority", "Medium"))

    for col_idx in range(1, LAST_COL + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.border = THIN_BORDER
        cell.font = FONT_DEFAULT
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        elif col_idx == 7:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = _priority_fill(tc.get("priority", ""))
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    max_lines = 1
    for f in ["test_item", "pre_condition", "data_test", "step", "expected"]:
        t = tc.get(f, "") or ""
        lines = t.count("\n") + max(1, len(t) // 50)
        max_lines = max(max_lines, lines)
    ws.row_dimensions[row].height = max(22, min(max_lines * 16, 280))


def create_mixed_workbook(output_path, module_name, module_code, sections, creator=""):
    """
    sections support 2 style mix trong 1 sheet:
    
    Style 1 - field_validation (có sub-section per field):
    {
        "name": "II. API Authentication",
        "field_groups": [
            {"field": "request_id", "required": "required", "test_cases": [...]},
            ...
        ]
    }
    
    Style 2 - plain (không có field grouping):
    {
        "name": "III. Happy path flow",
        "test_cases": [...]
    }
    """
    wb = Workbook()
    ws = wb.active
    ws.title = module_name[:31]

    _set_header_block(ws, module_code, creator)
    _set_column_widths(ws)

    current_row = 8
    for section in sections:
        _add_section_header(ws, current_row, section["name"])
        current_row += 1

        # Auto-detect style
        if "field_groups" in section:
            for fg in section["field_groups"]:
                _add_field_header(ws, current_row, fg["field"], fg.get("required", ""))
                current_row += 1
                for tc in fg.get("test_cases", []):
                    _add_test_case_row(ws, current_row, tc)
                    current_row += 1
        elif "test_cases" in section:
            for tc in section["test_cases"]:
                _add_test_case_row(ws, current_row, tc)
                current_row += 1

    ws.freeze_panes = "A8"
    wb.save(output_path)
    return output_path
